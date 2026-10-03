"""
Extraction pipeline tests.

Creates real sample files and runs end-to-end extraction without needing Qt
or a running application.  All processing is local/offline.
"""
from __future__ import annotations

import io
import tempfile
from pathlib import Path

import pytest
from PIL import Image, ImageDraw

from app.extraction import (
    DOCXExtractor,
    ExtractionError,
    ExtractionManager,
    ExtractionResult,
    ImageOCR,
    InputDetector,
    PDFTextExtractor,
    SVGExtractor,
    TextFileWriter,
    TextProcessor,
    UnsupportedInputError,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _white_image(text: str, size: tuple = (600, 120)) -> Image.Image:
    img = Image.new("RGB", size, "white")
    ImageDraw.Draw(img).text((20, 30), text, fill="black")
    return img


def _save_png(tmp_path: Path, text: str, name: str = "test.png") -> Path:
    path = tmp_path / name
    _white_image(text).save(path)
    return path


def _save_text_pdf(tmp_path: Path) -> Path:
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 100), "Direct PDF text line one.", fontsize=14)
    page.insert_text((72, 130), "Direct PDF text line two.", fontsize=14)
    path = tmp_path / "text.pdf"
    doc.save(str(path))
    doc.close()
    return path


def _save_scanned_pdf(tmp_path: Path) -> Path:
    import pymupdf
    img = _white_image("Scanned PDF OCR text", (700, 140))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_image(pymupdf.Rect(0, 0, 612, 200), stream=buf.read())
    path = tmp_path / "scanned.pdf"
    doc.save(str(path))
    doc.close()
    return path


def _save_mixed_pdf(tmp_path: Path) -> Path:
    import pymupdf
    img = _white_image("Scanned page two content", (700, 140))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    doc = pymupdf.open()
    p1 = doc.new_page(width=612, height=792)
    p1.insert_text((72, 100), "Page one direct text", fontsize=14)
    p2 = doc.new_page(width=612, height=792)
    p2.insert_image(pymupdf.Rect(0, 0, 612, 200), stream=buf.read())
    path = tmp_path / "mixed.pdf"
    doc.save(str(path))
    doc.close()
    return path


def _save_docx(tmp_path: Path) -> Path:
    from docx import Document as DocxDoc
    doc = DocxDoc()
    doc.add_heading("Test Heading", level=1)
    doc.add_paragraph("Test paragraph content here.")
    tbl = doc.add_table(rows=1, cols=2)
    tbl.rows[0].cells[0].text = "CellA"
    tbl.rows[0].cells[1].text = "CellB"
    path = tmp_path / "test.docx"
    doc.save(path)
    return path


def _save_svg_text(tmp_path: Path) -> Path:
    path = tmp_path / "vector.svg"
    path.write_bytes(
        b'<?xml version="1.0"?>'
        b'<svg xmlns="http://www.w3.org/2000/svg" width="300" height="80">'
        b'<text x="10" y="50" font-size="18">SVG Vector Text</text>'
        b'</svg>'
    )
    return path


def _save_svg_shapes(tmp_path: Path) -> Path:
    path = tmp_path / "shapes.svg"
    path.write_bytes(
        b'<?xml version="1.0"?>'
        b'<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100">'
        b'<rect x="0" y="0" width="200" height="100" fill="white"/>'
        b'<circle cx="100" cy="50" r="40" fill="blue"/>'
        b'</svg>'
    )
    return path


# ---------------------------------------------------------------------------
# InputDetector
# ---------------------------------------------------------------------------

class TestInputDetector:
    def test_detects_png(self, tmp_path):
        p = _save_png(tmp_path, "hello")
        assert InputDetector().detect(p).kind == "raster"

    def test_detects_jpg(self, tmp_path):
        p = tmp_path / "img.jpg"
        _white_image("hello").convert("RGB").save(p, format="JPEG")
        assert InputDetector().detect(p).kind == "raster"

    def test_detects_svg_text(self, tmp_path):
        p = _save_svg_text(tmp_path)
        assert InputDetector().detect(p).kind == "svg"

    def test_detects_pdf_by_magic(self, tmp_path):
        p = _save_text_pdf(tmp_path)
        assert InputDetector().detect(p).kind == "pdf"

    def test_detects_docx(self, tmp_path):
        p = _save_docx(tmp_path)
        assert InputDetector().detect(p).kind == "docx"

    def test_rejects_unsupported(self, tmp_path):
        p = tmp_path / "file.xyz"
        p.write_bytes(b"garbage")
        with pytest.raises(UnsupportedInputError):
            InputDetector().detect(p)

    def test_rejects_missing_file(self, tmp_path):
        with pytest.raises(ExtractionError):
            InputDetector().detect(tmp_path / "ghost.png")

    def test_rejects_corrupt_pdf(self, tmp_path):
        p = tmp_path / "fake.pdf"
        p.write_bytes(b"not a pdf at all")
        with pytest.raises(UnsupportedInputError):
            InputDetector().detect(p)


# ---------------------------------------------------------------------------
# TextProcessor
# ---------------------------------------------------------------------------

class TestTextProcessor:
    def test_strips_blank_edges(self):
        assert TextProcessor().normalize("\n\nHello\n\n") == "Hello"

    def test_collapses_multiple_blanks(self):
        assert TextProcessor().normalize("A\n\n\n\nB") == "A\n\nB"

    def test_removes_trailing_spaces(self):
        assert TextProcessor().normalize("Line   \nTwo  ") == "Line\nTwo"

    def test_normalises_crlf(self):
        assert TextProcessor().normalize("A\r\nB") == "A\nB"


# ---------------------------------------------------------------------------
# TextFileWriter
# ---------------------------------------------------------------------------

class TestTextFileWriter:
    def test_writes_beside_source(self, tmp_path):
        src = tmp_path / "doc.pdf"
        src.write_bytes(b"dummy")
        out = TextFileWriter().write(src, "hello")
        assert out == tmp_path / "doc.txt"
        assert out.read_text(encoding="utf-8").strip() == "hello"

    def test_avoids_overwrite(self, tmp_path):
        src = tmp_path / "doc.pdf"
        src.write_bytes(b"dummy")
        existing = tmp_path / "doc.txt"
        existing.write_text("original", encoding="utf-8")
        out = TextFileWriter().write(src, "new")
        assert out != existing
        assert existing.read_text(encoding="utf-8") == "original"


# ---------------------------------------------------------------------------
# SVGExtractor
# ---------------------------------------------------------------------------

class TestSVGExtractor:
    def test_extracts_vector_text(self, tmp_path):
        p = _save_svg_text(tmp_path)
        text = SVGExtractor(ImageOCR()).extract(p)
        assert "SVG Vector Text" in text

    def test_handles_shapes_without_crash(self, tmp_path):
        p = _save_svg_shapes(tmp_path)
        text = SVGExtractor(ImageOCR()).extract(p)
        assert isinstance(text, str)  # may be empty for pure shapes

    def test_rejects_malformed_svg(self, tmp_path):
        p = tmp_path / "bad.svg"
        p.write_bytes(b"<svg><not valid xml")
        with pytest.raises(ExtractionError):
            SVGExtractor(ImageOCR()).extract(p)


# ---------------------------------------------------------------------------
# DOCXExtractor
# ---------------------------------------------------------------------------

class TestDOCXExtractor:
    def test_extracts_paragraph(self, tmp_path):
        p = _save_docx(tmp_path)
        text = DOCXExtractor().extract(p)
        assert "Test paragraph content" in text

    def test_extracts_heading(self, tmp_path):
        p = _save_docx(tmp_path)
        text = DOCXExtractor().extract(p)
        assert "Test Heading" in text

    def test_extracts_table_cells(self, tmp_path):
        p = _save_docx(tmp_path)
        text = DOCXExtractor().extract(p)
        assert "CellA" in text
        assert "CellB" in text

    def test_handles_corrupt_docx(self, tmp_path):
        p = tmp_path / "bad.docx"
        p.write_bytes(b"PK fake zip garbage")
        with pytest.raises(ExtractionError):
            DOCXExtractor().extract(p)


# ---------------------------------------------------------------------------
# PDFTextExtractor
# ---------------------------------------------------------------------------

class TestPDFTextExtractor:
    def test_extracts_direct_text(self, tmp_path):
        p = _save_text_pdf(tmp_path)
        text = PDFTextExtractor(ImageOCR()).extract(p)
        assert "Direct PDF text" in text

    def test_includes_page_labels(self, tmp_path):
        p = _save_text_pdf(tmp_path)
        text = PDFTextExtractor(ImageOCR()).extract(p)
        assert "Page 1" in text

    def test_scanned_pdf_returns_string(self, tmp_path):
        p = _save_scanned_pdf(tmp_path)
        text = PDFTextExtractor(ImageOCR()).extract(p)
        assert isinstance(text, str)

    def test_mixed_pdf_has_both_pages(self, tmp_path):
        p = _save_mixed_pdf(tmp_path)
        text = PDFTextExtractor(ImageOCR()).extract(p)
        assert "Page 1" in text
        assert "Page 2" in text
        assert "Page one direct text" in text


# ---------------------------------------------------------------------------
# ExtractionManager end-to-end
# ---------------------------------------------------------------------------

class TestExtractionManager:
    def test_png_produces_txt_file(self, tmp_path):
        p = _save_png(tmp_path, "Hello extraction")
        try:
            result = ExtractionManager().extract_file(p)
            assert isinstance(result, ExtractionResult)
            assert result.output_path.exists()
            assert result.input_kind == "raster"
        except ExtractionError:
            pytest.skip("OCR returned no text from test image in this environment")

    def test_text_pdf_end_to_end(self, tmp_path):
        p = _save_text_pdf(tmp_path)
        result = ExtractionManager().extract_file(p)
        assert result.input_kind == "pdf"
        assert result.output_path.exists()
        assert "Direct PDF text" in result.text

    def test_docx_end_to_end(self, tmp_path):
        p = _save_docx(tmp_path)
        result = ExtractionManager().extract_file(p)
        assert result.input_kind == "docx"
        assert "Test paragraph" in result.text
        assert result.output_path.name.startswith("test")

    def test_svg_vector_end_to_end(self, tmp_path):
        p = _save_svg_text(tmp_path)
        result = ExtractionManager().extract_file(p)
        assert result.input_kind == "svg"
        assert "SVG Vector Text" in result.text

    def test_unsupported_raises(self, tmp_path):
        p = tmp_path / "vid.mp4"
        p.write_bytes(b"fake mp4 data")
        with pytest.raises(UnsupportedInputError):
            ExtractionManager().extract_file(p)

    def test_cancellation_raises(self, tmp_path):
        p = _save_text_pdf(tmp_path)
        with pytest.raises(ExtractionError, match="cancelled"):
            ExtractionManager().extract_file(p, cancelled=lambda: True)


# ---------------------------------------------------------------------------
# Worker process
# ---------------------------------------------------------------------------

class TestExtractionWorker:
    def test_bad_path_returns_error(self):
        from app.extraction_worker import start_extraction_worker
        proc, queue, cancel = start_extraction_worker("/no/such/file.png")
        proc.join(timeout=30)
        messages = []
        while not queue.empty():
            messages.append(queue.get_nowait())
        result_msgs = [m for m in messages if m[0] == "result"]
        assert len(result_msgs) == 1
        assert result_msgs[0][1] == "error"

    def test_terminate_does_not_hang(self, tmp_path):
        from app.extraction_worker import start_extraction_worker
        p = _save_text_pdf(tmp_path)
        proc, queue, cancel = start_extraction_worker(str(p))
        proc.terminate()
        proc.join(timeout=5)
        assert not proc.is_alive()

    def test_soft_cancel_event(self, tmp_path):
        from app.extraction_worker import start_extraction_worker
        p = _save_text_pdf(tmp_path)
        proc, queue, cancel = start_extraction_worker(str(p))
        cancel.set()   # signal before the worker even starts OCR
        proc.join(timeout=30)
        assert not proc.is_alive()
